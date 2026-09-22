#!/usr/bin/env python3
"""
Build the Test Cases / Traceability Matrix / Validation Summary workbook.

Usage:
    python3 build_workbook.py --input payload.json --output Test_Cases_Feature.xlsx

Payload schema (see SKILL.md for the full description and an example):
{
  "test_cases": [
    {
      "id": "TC-001", "user_story": "...", "acceptance_criteria": "...",
      "title": "...", "preconditions": "...", "test_data": "...", "steps": "...",
      "expected_results": "...", "test_type": "Functional", "testing_type": "UI",
      "priority": "High", "severity": "High"
    }
  ],
  "traceability": [
    {"user_story": "...", "acceptance_criteria": "AC-01", "test_case_ids": "TC-001",
     "coverage_status": "Covered", "coverage_type": "Positive", "notes": ""}
  ],
  "validation_summary": [
    {"item": "User stories identified", "result": "1", "details": "..."}
  ]
}

"Actual Results" is not part of the input schema -- this script always writes it blank, since the
skill's rules require that column to never contain fabricated execution results.
"""

import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

TEST_CASE_COLUMNS = [
    ("Test Case ID", "id", 12),
    ("User Stories", "user_story", 32),
    ("Acceptance Criteria", "acceptance_criteria", 32),
    ("Title", "title", 30),
    ("Preconditions", "preconditions", 26),
    ("Test Data", "test_data", 26),
    ("Steps", "steps", 36),
    ("Expected Results", "expected_results", 34),
    ("Test Type", "test_type", 16),
    ("TestingType", "testing_type", 14),
    ("Actual Results", None, 16),  # always blank
    ("Priority", "priority", 12),
    ("Severity", "severity", 12),
]

TRACEABILITY_COLUMNS = [
    ("User Story", "user_story", 30),
    ("Acceptance Criteria", "acceptance_criteria", 20),
    ("Test Case IDs", "test_case_ids", 24),
    ("Coverage Status", "coverage_status", 18),
    ("Coverage Type", "coverage_type", 18),
    ("Notes", "notes", 30),
]

VALIDATION_COLUMNS = [
    ("Validation Item", "item", 34),
    ("Result", "result", 16),
    ("Details", "details", 44),
]

ALLOWED_PRIORITY_SEVERITY = {"Critical", "High", "Medium", "Low"}
ALLOWED_COVERAGE_STATUS = {"Covered", "Partially Covered", "Not Covered", "Needs Clarification"}


def write_sheet(ws: Worksheet, columns, rows):
    headers = [c[0] for c in columns]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"

    for row in rows:
        values = []
        for header, key, _ in columns:
            if key is None:
                values.append("")  # Actual Results, always blank
            else:
                values.append(row.get(key, ""))
        ws.append(values)

    last_row = ws.max_row
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    for col_idx, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for r in range(2, last_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = WRAP_TOP


def validate_payload(payload):
    warnings = []
    seen_ids = set()
    for tc in payload.get("test_cases", []):
        tc_id = tc.get("id", "")
        if tc_id in seen_ids:
            warnings.append(f"Duplicate Test Case ID detected: {tc_id}")
        seen_ids.add(tc_id)
        for field in ("priority", "severity"):
            val = tc.get(field, "")
            if val and val not in ALLOWED_PRIORITY_SEVERITY:
                warnings.append(f"{tc_id}: invalid {field} value '{val}'")
        if not tc.get("steps"):
            warnings.append(f"{tc_id}: missing Steps")
        if not tc.get("expected_results"):
            warnings.append(f"{tc_id}: missing Expected Results")

    for tr in payload.get("traceability", []):
        status = tr.get("coverage_status", "")
        if status and status not in ALLOWED_COVERAGE_STATUS:
            warnings.append(
                f"Traceability row for {tr.get('acceptance_criteria', '?')}: "
                f"invalid coverage_status '{status}'"
            )
    return warnings


def build_workbook(payload, output_path):
    wb = Workbook()

    ws_tc = wb.active
    ws_tc.title = "Test Cases"
    write_sheet(ws_tc, TEST_CASE_COLUMNS, payload.get("test_cases", []))

    ws_tm = wb.create_sheet("Traceability Matrix")
    write_sheet(ws_tm, TRACEABILITY_COLUMNS, payload.get("traceability", []))

    ws_vs = wb.create_sheet("Validation Summary")
    write_sheet(ws_vs, VALIDATION_COLUMNS, payload.get("validation_summary", []))

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to the JSON payload")
    parser.add_argument("--output", required=True, help="Output .xlsx path")
    args = parser.parse_args()

    with open(args.input, "r", encoding="utf-8") as f:
        payload = json.load(f)

    warnings = validate_payload(payload)
    if warnings:
        print("Quality gate warnings (fix these in the payload before relying on the sheet):",
              file=sys.stderr)
        for w in warnings:
            print(f"  - {w}", file=sys.stderr)

    build_workbook(payload, args.output)
    print(f"Wrote {args.output}")
    if warnings:
        sys.exit(1)


if __name__ == "__main__":
    main()
